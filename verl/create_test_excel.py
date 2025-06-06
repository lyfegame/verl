#!/usr/bin/env python3
"""Create test Excel files for testing the Excel dataset."""

import openpyxl
from openpyxl import Workbook
import os

def create_test_excel_files():
    """Create test Excel files."""
    
    # Create directory if it doesn't exist
    os.makedirs("/tmp", exist_ok=True)
    
    # Create initial Excel file
    wb_initial = Workbook()
    ws = wb_initial.active
    ws.title = "Payments"
    
    # Add headers
    headers = ["Invoice ID", "Customer", "Due Date", "Amount", "Status"]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    
    # Add some data
    data = [
        ["INV001", "Acme Corp", "2023-01-15", 1000, "Paid"],
        ["INV002", "Tech Inc", "2023-01-20", 1500, "Paid"],
        ["INV003", "Global Ltd", "2023-01-25", 2000, "Unpaid"],
        ["INV004", "Smart Co", "2023-02-01", 1200, "Paid"],
        ["INV005", "Future LLC", "2023-02-05", 1800, "Unpaid"],
    ]
    
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # Save initial file
    wb_initial.save("/tmp/test_initial.xlsx")
    print("Created /tmp/test_initial.xlsx")
    
    # Create reference file (with the expected changes)
    wb_ref = openpyxl.load_workbook("/tmp/test_initial.xlsx")
    ws_ref = wb_ref.active
    
    # Add "Days Overdue" header
    ws_ref.cell(row=1, column=6, value="Days Overdue")
    
    # Add formulas for unpaid items
    for row in range(2, 7):
        status = ws_ref.cell(row=row, column=5).value
        if status == "Unpaid":
            # Add formula (simplified for testing)
            ws_ref.cell(row=row, column=6, value=f'=TODAY()-C{row}')
        else:
            ws_ref.cell(row=row, column=6, value="")
    
    # Add conditional formatting
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill
    
    red_fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
    ws_ref.conditional_formatting.add('F2:F6',
        CellIsRule(operator='greaterThan', formula=['0'], fill=red_fill))
    
    # Freeze top row
    ws_ref.freeze_panes = 'A2'
    
    # Save reference file
    wb_ref.save("/tmp/test_reference.xlsx")
    print("Created /tmp/test_reference.xlsx")
    
    # Create a final file (same as initial for now - in real scenario, 
    # this would be generated based on model output)
    wb_initial.save("/tmp/test_final.xlsx")
    print("Created /tmp/test_final.xlsx")
    
    print("\nTest Excel files created successfully!")

if __name__ == "__main__":
    create_test_excel_files() 